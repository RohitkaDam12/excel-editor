"""
main.py - Main entry point for the AI-Powered Excel Editor Prototype.

This application demonstrates safe AI-driven Excel editing:
  1. The user provides an Excel file and a natural language instruction.
  2. The LLM parses the instruction into structured JSON (never generating code).
  3. The Python Action Router validates and executes predefined openpyxl functions.
  4. The modified workbook is saved to output.xlsx.
"""

import argparse
import json
import os
import sys
import openpyxl

# Ensure UTF-8 output encoding on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from excel_operations import execute_operation, inspect_workbook
from llm import understand_request


def save_workbook_safely(wb: openpyxl.Workbook, output_path: str) -> str:
    """
    Saves the workbook safely. If output_path is locked by another application (e.g. Excel),
    automatically saves to an alternative filename (e.g. output_1.xlsx) so work is never lost.
    """
    try:
        wb.save(output_path)
        return output_path
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        counter = 1
        while counter <= 100:
            alt_path = f"{base}_{counter}{ext}"
            try:
                wb.save(alt_path)
                print(f"\n[Notice] '{output_path}' is currently open and locked in Microsoft Excel.")
                print(f"Saved to alternative file instead: {alt_path}")
                return alt_path
            except PermissionError:
                counter += 1
        raise PermissionError(f"Could not save workbook: all files from {output_path} to {alt_path} are locked.")


def process_excel_request(
    file_path: str,
    user_instruction: str,
    output_path: str = "output.xlsx"
) -> bool:
    """
    Executes the full pipeline:
      User Request -> LLM JSON -> Python Action -> OpenPyXL Operation -> Output File
    """
    if not os.path.exists(file_path):
        print(f"\nError: File '{file_path}' does not exist.")
        return False

    print("\n" + "=" * 50)
    print("USER REQUEST")
    print("=" * 50)
    print(f"Instruction: {user_instruction}")
    print(f"Input file:  {file_path}")

    # 1. Load Workbook
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return False

    # 2. Inspect Workbook (Sheet name, rows, columns, headers)
    inspection_info = inspect_workbook(ws)

    # 3. LLM Parsing
    print("\n|")
    print("v")
    print("LLM JSON")
    print("-" * 50)
    try:
        structured_json = understand_request(
            user_prompt=user_instruction,
            headers=inspection_info.get("headers")
        )
    except Exception as e:
        print(f"Error parsing instruction: {e}")
        return False

    print("LLM generated operations:")
    print(json.dumps(structured_json, indent=2))

    operations = structured_json.get("operations", [])
    if not operations:
        print("No operations returned by LLM.")
        return False

    # 4. Python Action Router & OpenPyXL Operations
    print("\n|")
    print("v")
    print("PYTHON ACTION & OPENPYXL OPERATION")
    print("-" * 50)
    success_count = 0
    exported_files = []
    for op in operations:
        action_name = op.get("action")
        print(f"\nRouter dispatching: {action_name}")
        op_success = execute_operation(ws, op, wb=wb)
        if op_success:
            success_count += 1
            if action_name == "export_to_file":
                out_f = op.get("output_file") or op.get("file") or op.get("filename")
                if out_f:
                    if not str(out_f).strip().lower().endswith(".xlsx"):
                        out_f = str(out_f).strip() + ".xlsx"
                    exported_files.append(out_f)

    # 5. Save Modified Workbook
    print("\n|")
    print("v")
    print("OUTPUT FILE")
    print("-" * 50)
    if success_count > 0:
        if exported_files:
            print("\n" + "=" * 60)
            print("  🎉 NEW EXCEL FILE(S) EXPORTED SUCCESSFULLY:")
            for ef in exported_files:
                abs_ef = os.path.abspath(ef)
                print(f"  -> File: {ef}")
                print(f"     Full path: {abs_ef}")
            print("=" * 60)

        only_exports = all(op.get("action") == "export_to_file" for op in operations)
        if only_exports:
            print(f"\n[NOTE] Your filtered data was saved directly into the new file above.")
            print(f"The input file '{file_path}' was left unmodified.")
            if output_path and output_path != "output.xlsx":
                saved_path = save_workbook_safely(wb, output_path)
                print(f"Saved copy of source to: {saved_path}")
        else:
            saved_path = save_workbook_safely(wb, output_path)
            print(f"\nSaved modified workbook to:\n{saved_path}")
            if len(wb.sheetnames) > 1:
                active_name = wb.active.title if wb.active else wb.sheetnames[0]
                print(f"Sheets in modified workbook: {wb.sheetnames} (Active: '{active_name}')")
            print(f"Successfully applied {success_count}/{len(operations)} operation(s).")
        return True
    else:
        print("\nNo changes saved because all operations failed validation or execution.")
        return False


def run_cli() -> None:
    """Interactive CLI mode."""
    parser = argparse.ArgumentParser(description="AI-Powered Excel Editor Prototype")
    parser.add_argument("--file", "-f", help="Path to input .xlsx file", default=None)
    parser.add_argument("--instruction", "-i", help="Natural-language instruction", default=None)
    parser.add_argument("--output", "-o", help="Path to output .xlsx file", default="output.xlsx")
    args = parser.parse_args()

    # Get file path
    if args.file:
        file_path = args.file
    else:
        print("=" * 45)
        print("     AI-Powered Excel Editor Prototype       ")
        print("=" * 45)
        file_input = input("Enter Excel file path [default: input.xlsx]: ").strip()
        file_path = file_input if file_input else "input.xlsx"

    # Get instruction
    if args.instruction:
        user_instruction = args.instruction
    else:
        user_instruction = input("\nWhat do you want to change?\n> ").strip()

    if not user_instruction:
        print("No instruction provided. Exiting.")
        sys.exit(1)

    process_excel_request(file_path, user_instruction, args.output)


if __name__ == "__main__":
    run_cli()
